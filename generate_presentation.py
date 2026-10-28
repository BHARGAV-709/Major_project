import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme colors
    PRIMARY = RGBColor(15, 23, 42)      # Slate 900
    ACCENT = RGBColor(37, 99, 235)      # Royal Blue
    ACCENT_LIGHT = RGBColor(239, 246, 255) # Blue 50
    CARD_BG = RGBColor(248, 250, 252)   # Slate 50
    TEXT_DARK = RGBColor(30, 41, 59)    # Slate 800
    TEXT_MUTED = RGBColor(100, 116, 139)# Slate 500
    BORDER_COL = RGBColor(226, 232, 240)# Slate 200
    WHITE = RGBColor(255, 255, 255)
    GREEN = RGBColor(16, 185, 129)
    AMBER = RGBColor(245, 158, 11)

    def add_header(slide, title_text, category_text="ACADEMIC PROJECT PRESENTATION"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT
        p0.space_after = Pt(4)
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY

    def add_card(slide, left, top, width, height, title, items, badge=""):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = BORDER_COL
        shape.line.width = Pt(1)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = f"{title}   [{badge}]" if badge else title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT
        p.space_after = Pt(10)

        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = f"•  {item}"
            p_item.font.size = Pt(11)
            p_item.font.color.rgb = TEXT_DARK
            p_item.space_after = Pt(6)

    def set_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PRIMARY
    bg1.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "MAJOR PROJECT DEFENSE & PRESENTATION"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = ACCENT
    p_badge.space_after = Pt(12)

    p_title = tf1.add_paragraph()
    p_title.text = "Resource Allocation and Scheduling Optimization for Enterprise Analytics"
    p_title.font.size = Pt(28)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.space_after = Pt(8)

    p_sub = tf1.add_paragraph()
    p_sub.text = "A Comparative Study of Heuristic, Metaheuristic, and Reinforcement Learning Methods"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)
    p_sub.space_after = Pt(36)

    p_meta = tf1.add_paragraph()
    p_meta.text = "Project Team: [Student Name(s) / Roll No.]\nDepartment of Computer Science & Engineering\n[Institution / University Name]  |  Academic Year 2026"
    p_meta.font.size = Pt(13)
    p_meta.font.color.rgb = RGBColor(203, 213, 225)

    set_notes(s1, "Good morning respected committee members, examiners, and faculty. Today I present our major project: Resource Allocation and Scheduling Optimization for Enterprise Analytics. This research investigates multi-resource scheduling across 10 distinct algorithms spanning classical heuristics, stochastic metaheuristics, and reinforcement learning within a discrete-event simulation framework.")

    # ==========================================
    # SLIDE 2: Introduction & Motivation
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Introduction: Scheduling Challenges in Enterprise Analytics")
    add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Modern Cloud Analytics Reality", [
        "Analytics clusters (Kubernetes, YARN, Slurm) process heterogeneous, multi-tenant workloads.",
        "Workloads combine batch ETL, deep learning training, and streaming data analytics.",
        "Jobs request multi-dimensional resource vectors simultaneously: CPU cores, RAM, and GPU accelerators.",
        "High infrastructure costs: Cloud compute OpEx requires strict resource utilization efficiency."
    ], "Context")
    add_card(s2, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0), "The Core Operational Challenges", [
        "Multi-Resource Stranding: GPU units sit idle when host CPU or RAM is exhausted.",
        "SLA Violation Risk: Missing negotiated completion deadlines incurs severe monetary penalties.",
        "Unpredictable Queue Dynamics: High contention and bursty arrivals cause massive waiting delays.",
        "Need for Simulation: Evaluating algorithms directly on physical clusters is expensive, noisy, and non-reproducible."
    ], "Motivation")
    set_notes(s2, "In enterprise environments, modern workloads are heterogeneous and multi-dimensional. A job cannot run if it secures 8 CPU cores but lacks 16 GB of RAM or 1 GPU accelerator. This causes resource stranding. Furthermore, testing scheduling algorithms directly on live production clusters is volatile and expensive. Discrete-event simulation allows us to benchmark algorithms fairly under identical conditions.")

    # ==========================================
    # SLIDE 3: Problem Statement
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Problem Statement: Multi-Objective NP-Hard Allocation")
    add_card(s3, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), "System Formulation", [
        "Aggregated cluster resources: Total CPU (C_cpu), Memory (C_mem), and GPU (C_gpu).",
        "Set of N jobs J = {j1, j2, ... jN}.",
        "Each job specifies: arrival, runtime, CPU cores, RAM MB, GPU units, priority, deadline, budget.",
        "Non-preemptive execution: once started, a job runs uninterrupted until completion."
    ], "Model")
    add_card(s3, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), "Physical Constraints", [
        "Strict Capacity Limits: Concurrent consumption of running jobs must never exceed cluster capacity at any time t.",
        "Causal Precedence: No job may start before its arrival timestamp (t_start >= t_arrival).",
        "Permanent Infeasibility: Jobs exceeding total cluster capacity must be rejected upfront without deadlocking queues."
    ], "Constraints")
    add_card(s3, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), "Conflicting Objectives", [
        "Minimizing Makespan: Clear full workload as rapidly as possible.",
        "Minimizing Waiting Time: Reduce tenant queue latency.",
        "Maximizing Utilization: High continuous duty cycle on CPU/RAM/GPU.",
        "Minimizing SLA Breaches & Cost: Avoid SLA penalty fees ($25/breach)."
    ], "Trade-offs")
    set_notes(s3, "Mathematically, multi-resource job scheduling is NP-hard. We have three hard constraints: capacity limits across CPU, memory, and GPU; causal precedence where jobs cannot execute before arrival; and non-preemption. Our objectives inherently conflict: maximizing packing density can delay urgent tasks, while strictly prioritizing deadlines can fragment resource allocations.")

    # ==========================================
    # SLIDE 4: Objectives & Scope
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Project Aim, Objectives, and Technical Scope")
    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Measurable Project Objectives", [
        "Build Discrete-Event Simulator: Implement continuous-time simulator with min-heap priority event queue.",
        "Implement 10 Algorithms: Cover classical heuristics, global metaheuristics, and reinforcement learning.",
        "Data Preprocessing & Workload Generation: Support Google cluster traces (10K tasks) and synthetic scenarios.",
        "Standardized Benchmarking: Establish fair evaluation with master seeds and 5-seed stochastic testing.",
        "Interactive Dashboard: Deliver React/Vite UI with Gantt charts, sequential order flow, and CSV export."
    ], "Deliverables")
    add_card(s4, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0), "Project Scope & Verified Boundaries", [
        "Simulation Context: All benchmarks run within the discrete-event simulation engine, not a live physical cluster.",
        "Hardware Model: Standardized cluster of 40 CPU cores, 256 GB RAM, and 8 GPUs.",
        "Non-Preemptive Execution: Tasks run continuously until finish; no checkpoint/migration modeling.",
        "Synchronous REST Architecture: Fast browser feedback (<10s) with compact RL training budgets (25–60 episodes).",
        "Upfront Rejection: Permanently infeasible jobs marked terminal REJECTED; metrics evaluate executed jobs."
    ], "Boundaries")
    set_notes(s4, "Our goal was to build a complete, verifiable research and benchmarking platform. We specifically implemented 10 algorithms, standardized the simulation environment on a 40-core, 256GB RAM, 8-GPU cluster, and built an interactive dashboard. We are clear about our boundaries: this is a discrete-event simulation engine, scheduling is non-preemptive, and rejected jobs are reported cleanly without corrupting runtime metrics.")

    # ==========================================
    # SLIDE 5: Existing Approaches & Motivation
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Existing Approaches & Limitations of Single Policies")
    add_card(s5, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), "Classical Heuristics", [
        "FCFS: Zero overhead but suffers from convoy effect; long tasks block short jobs.",
        "SJF: Minimizes wait time in simple queues, but starves long tasks and ignores deadlines.",
        "Priority: Protects high-tier tasks but causes starvation for lower-priority jobs.",
        "EDF: Effective in uniprocessor real-time systems, but causes domino SLA breaches under multi-resource overload.",
        "Greedy Best-Fit: High instant packing, but blind to future arrivals."
    ], "Heuristics")
    add_card(s5, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), "Metaheuristic Search", [
        "Explores global combinatorial sequence permutations.",
        "GA: Employs crossover and mutation on permutation chromosomes.",
        "PSO: Maps continuous particle velocities to discrete job rankings.",
        "SA: Escapes local minima via Boltzmann probabilistic uphill acceptance.",
        "Limitation: Higher computational complexity than simple one-pass heuristics."
    ], "Metaheuristics")
    add_card(s5, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), "Reinforcement Learning", [
        "Learns dispatch policies dynamically through environment interaction.",
        "Tabular Q-Learning: Temporal difference updates over state buckets.",
        "DQN: PyTorch neural network with experience replay buffer.",
        "Project Scope: Uses a compact training budget (25–30 episodes); does not establish asymptotic convergence or global optimality."
    ], "RL Methods")
    set_notes(s5, "Why not use just one algorithm? Real-world heuristics have severe blind spots: FCFS suffers from convoy delays, SJF delays long batch jobs under heavy queues, and Greedy Best-Fit creates fragmented residual capacity. Metaheuristics optimize globally across permutations. RL agents offer adaptive formulation, but under our compact training budget of 25 to 30 episodes, they do not establish convergence or outperform metaheuristics.")

    # ==========================================
    # SLIDE 6: Proposed System & Pipeline
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Proposed System: End-to-End Workflow Pipeline")
    add_card(s6, Inches(0.8), Inches(1.8), Inches(2.7), Inches(5.0), "1. Workload Source", [
        "Google Trace CSV (10K raw cluster tasks).",
        "Synthetic Scenarios: LIGHT, MEDIUM, HEAVY, BURSTY, OVERLOADED.",
        "Reproducible Master Random Seed (Seed 42)."
    ], "Input")
    add_card(s6, Inches(3.8), Inches(1.8), Inches(2.7), Inches(5.0), "2. Preprocessing & Filter", [
        "8-step cleaning pipeline.",
        "Missing value mean imputation.",
        "IQR outlier clipping & min-max scaling.",
        "Admissibility Filter: Checks if job exceeds cluster capacity. Rejects permanently infeasible jobs."
    ], "Preparation")
    add_card(s6, Inches(6.8), Inches(1.8), Inches(2.7), Inches(5.0), "3. Engine & Scheduler", [
        "SchedulingDecisionEngine orchestrates algorithm execution.",
        "DiscreteEventSimulator processes ARRIVAL & COMPLETION events via min-heap queue.",
        "Multi-resource tracker updates CPU, RAM, and GPU capacity."
    ], "Simulation")
    add_card(s6, Inches(9.8), Inches(1.8), Inches(2.7), Inches(5.0), "4. Evaluation & UI", [
        "PerformanceEvaluator calculates 6 formal metrics.",
        "SQLite database logs run history.",
        "React frontend renders interactive Gantt timeline & CSV export."
    ], "Output")
    set_notes(s6, "This slide illustrates the full lifecycle: Workloads enter either from the Google cluster dataset or synthetic generator. They pass through our 8-step preprocessing pipeline and capacity filter. Valid jobs enter the discrete-event simulator where one of 10 algorithms controls dispatching. The PerformanceEvaluator computes standard metrics, results are persisted to SQLite, and the React frontend renders interactive Gantt timelines.")

    # ==========================================
    # SLIDE 7: System Architecture
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "System Architecture: Decoupled Multi-Tier Platform")
    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3), "Frontend Presentation Tier (React 18 + Vite)", [
        "Dashboard (/): System KPIs, algorithm comparisons, and cluster status.",
        "Simulation (/simulation): Single run execution, Gantt chart, status bars.",
        "Results (/results): Sequential execution order (#1, #2, #3...), CSV export.",
        "Built with Tailwind CSS, Recharts SVG charting, and Lucide Icons."
    ], "Client")
    add_card(s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.3), "Backend Application Tier (FastAPI + Python 3.10)", [
        "API Gateway (app/api/routes.py) with strict Pydantic validation.",
        "Configurable CORS origins (ports 3000, 5173).",
        "SQLite Relational DB via SQLAlchemy ORM (cluster_scheduler.db).",
        "Automated OpenAPI documentation at /docs."
    ], "Server")
    add_card(s7, Inches(0.8), Inches(4.4), Inches(11.6), Inches(2.4), "Core Simulation & Algorithmic Layer", [
        "SchedulingDecisionEngine: Unified orchestrator standardizing inputs and outputs across all 10 algorithms.",
        "DiscreteEventSimulator: Continuous-time engine managing min-heap priority event queue (O(log N)).",
        "ClusterResourceState: Synchronized multi-resource tracker for CPU cores, RAM MB, and GPU units.",
        "PerformanceEvaluator: Calculates Makespan, Waiting Time, Utilization, SLA Rate, Cost, and Composite Score."
    ], "Core Engine")
    set_notes(s7, "Our architecture strictly decouples presentation from core logic. The frontend is built with React 18, Vite, and Tailwind CSS. The backend utilizes FastAPI with Pydantic type safety and SQLite persistence. Crucially, the SchedulingDecisionEngine acts as a facade, ensuring every algorithm interfaces with the exact same DiscreteEventSimulator and PerformanceEvaluator.")

    # ==========================================
    # SLIDE 8: Ten Scheduling Algorithms
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Ten Implemented Algorithms Across Three Paradigms")
    add_card(s8, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), "Deterministic Heuristics (5)", [
        "1. FCFS: Dispatches by arrival timestamp t_arrival.",
        "2. SJF: Prioritizes shortest runtime t_runtime.",
        "3. Priority: Dispatches by static priority rank (1 to 5).",
        "4. EDF: Dispatches by earliest SLA deadline d_j.",
        "5. Greedy Best-Fit: Minimizes residual leftover cluster capacity across CPU, RAM, and GPU."
    ], "Classical")
    add_card(s8, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), "Global Metaheuristics (3)", [
        "6. Genetic Algorithm (GA): Permutation chromosomes, Tournament selection, Order Crossover (OX1), Swap mutation.",
        "7. Particle Swarm Optimization (PSO): Continuous particle velocities mapped to permutations via Smallest Position Value (SPV).",
        "8. Simulated Annealing (SA): 2-opt swaps with Metropolis probabilistic acceptance: P = exp(dF / T)."
    ], "Optimization")
    add_card(s8, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), "Reinforcement Learning (2)", [
        "9. Tabular Q-Learning: State discretization into 4-bin buckets; Bellman TD value updates: Q(s,a) <- Q(s,a) + a[r + g max Q - Q].",
        "10. Deep Q-Network (DQN): PyTorch MLP (6 -> 64 -> 64 -> |A|), Experience Replay Buffer (N=2000), Target Network sync.",
        "Post-Training Eval: Evaluated deterministically (eps = 0.0) across identical workloads."
    ], "Learning Agents")
    set_notes(s8, "We implemented exactly 10 algorithms across three broad families: Five deterministic heuristics for baseline efficiency; three metaheuristics that optimize global permutation schedules using evolutionary, swarm, and thermal physics models; and two reinforcement learning agents that learn dispatch policies through environmental rewards and neural network approximation.")

    # ==========================================
    # SLIDE 9: Evaluation Metrics & Formulations
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Formal Evaluation Metrics & Formulas")
    add_card(s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3), "1. Makespan & 2. Waiting Time", [
        "Makespan (M): Total simulation completion timestamp: M = max(t_finish_j).",
        "Average Waiting Time (W): Mean latency from arrival to execution start:",
        "W_bar = (1 / |J|) * sum(t_start_j - t_arrival_j).",
        "Reflects queue delay experienced by admitted jobs."
    ], "Time Metrics")
    add_card(s9, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.3), "3. Multi-Resource Utilization (U)", [
        "Dimension occupancy: U_dim = (sum demand_j * runtime_j) / (C_dim * M) * 100.",
        "Evaluated for CPU, RAM, and GPU dimensions independently.",
        "Overall Utilization: U_overall = 0.45*U_cpu + 0.45*U_mem + 0.10*U_gpu.",
        "Measures continuous active duty cycle across cluster hardware."
    ], "Hardware Efficiency")
    add_card(s9, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.4), "4. SLA Violation Rate & 5. Financial Cost", [
        "SLA Violation Rate: R_sla = (sum [t_finish_j > deadline_j] / |J|) * 100.",
        "Operating Cost Model: Hardware usage rates + SLA violation penalties:",
        "Cost = sum(cpu*0.00005) + sum(mem*0.00001) + sum(gpu*0.0002) + (Violations * $25.00).",
        "SLA penalties severely penalize overdue execution."
    ], "SLA & Cost")
    add_card(s9, Inches(6.8), Inches(4.4), Inches(5.6), Inches(2.4), "6. Composite Efficiency Score (0 to 100)", [
        "Unified academic evaluation index balancing all dimensions:",
        "Score = max(0, min(100, 0.40*U_overall + 0.40*(100 - R_sla) + 0.20*max(0, 100 - W_bar))).",
        "Waiting Saturation: When W_bar >= 100s under heavy load, waiting penalty saturates to 0, leaving utilization and SLA to drive score."
    ], "Composite Score")
    set_notes(s9, "All 10 algorithms are evaluated through the exact same formulas. Makespan measures completion time. Waiting time measures queue latency. Resource utilization measures multi-dimensional hardware occupancy. SLA rate tracks deadline breaches, while our cost model combines compute usage charges with a $25 breach penalty. Finally, our Composite Efficiency Score normalizes everything into a 0 to 100 scale.")

    # ==========================================
    # SLIDE 10: Experimental Setup & Fairness
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Experimental Setup & Scientific Rigor")
    add_card(s10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Standardized Simulation Environment", [
        "Cluster Capacity: 40 CPU cores, 256 GB RAM, 8 GPU accelerators.",
        "Workload Scale: Standardized at exactly N = 25 jobs per scenario.",
        "Reproducibility: Generated with deterministic Master Seed 42 and saved to benchmark_workloads_standard.json.",
        "Fairness Guarantee: Every algorithm receives an identical deep copy of the job list within each scenario; no shared memory mutation."
    ], "Cluster Setup")
    add_card(s10, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0), "Statistical Reporting Methodology", [
        "Deterministic Algorithms (FCFS, SJF, Priority, EDF, Greedy): Evaluated via single exact execution run.",
        "Stochastic Algorithms (GA, PSO, SA, Q-Learning, DQN): Evaluated across 5 independent random seeds (10, 20, 30, 40, 50).",
        "Reporting Format: Mean +/- Standard Deviation.",
        "Single Source of Truth: All results verified directly against standardized_benchmark_results.json."
    ], "Methodology")
    set_notes(s10, "Scientific fairness was paramount in our benchmarking. Every algorithm received an identical deep copy of the 25-job workload trace generated with master seed 42. For deterministic algorithms like FCFS and SJF, we report single runs. For stochastic optimizers and RL agents, we report the Mean and Standard Deviation across 5 independent random seeds. All data matches our standardized JSON artifact.")

    # ==========================================
    # SLIDE 11: Benchmark Results: Medium Contention
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Benchmark Results: Medium Contention Scenario (25 Jobs)")
    
    # Create Table on Slide
    table_shape = s11.shapes.add_table(11, 7, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    table = table_shape.table
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(1.8)
    table.columns[2].width = Inches(1.6)
    table.columns[3].width = Inches(1.5)
    table.columns[4].width = Inches(1.6)
    table.columns[5].width = Inches(1.6)
    table.columns[6].width = Inches(1.8)

    headers = ["Algorithm", "Type", "Makespan (s)", "Avg Wait (s)", "Util (%)", "SLA Viol (%)", "Score (0-100)"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = WHITE

    med_data = [
        ["PSO", "Stochastic", "176.45 ± 0.35", "8.95 ± 0.44", "44.21% ± 0.09%", "9.6% ± 2.0%", "72.06 ± 0.71 (1st)"],
        ["GA", "Stochastic", "175.92 ± 0.35", "8.57 ± 0.57", "44.35% ± 0.09%", "11.2% ± 1.6%", "71.54 ± 0.55 (2nd)"],
        ["SJF", "Deterministic", "178.18", "8.07", "43.78%", "12.0%", "71.10 (3rd)"],
        ["SA", "Stochastic", "179.01 ± 3.78", "8.68 ± 0.38", "43.60% ± 0.90%", "13.6% ± 3.2%", "70.26 ± 1.15"],
        ["Priority", "Deterministic", "189.61", "10.24", "41.14%", "20.0%", "66.41"],
        ["EDF", "Deterministic", "180.77", "9.25", "43.15%", "24.0%", "65.81"],
        ["Greedy", "Deterministic", "176.62", "9.87", "44.17%", "28.0%", "64.49"],
        ["FCFS", "Deterministic", "176.62", "10.21", "44.17%", "28.0%", "64.43"],
        ["DQN", "Stochastic", "207.66 ± 10.49", "19.19 ± 2.33", "37.66% ± 1.91%", "32.0% ± 5.1%", "58.43 ± 2.87"],
        ["Q-Learning", "Stochastic", "202.59 ± 12.06", "20.17 ± 2.62", "38.64% ± 2.28%", "33.6% ± 2.0%", "57.98 ± 0.74"]
    ]
    for row_idx, row_vals in enumerate(med_data):
        for col_idx, val in enumerate(row_vals):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = val
            cell.fill.solid()
            if row_idx == 0:
                cell.fill.fore_color.rgb = ACCENT_LIGHT
            else:
                cell.fill.fore_color.rgb = CARD_BG
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.color.rgb = TEXT_DARK
                if row_idx == 0:
                    p.font.bold = True

    set_notes(s11, "Under MEDIUM contention, measured benchmark results show PSO achieved the highest overall score of 72.06, followed closely by GA at 71.54. SJF performed well under this moderate load, achieving 71.10 by rapidly dispatching short jobs. FCFS and Greedy achieved moderate scores around 64.4, while DQN and Q-Learning scored ~58 under their compact training budget and summary state representation.")

    # ==========================================
    # SLIDE 12: Benchmark Results: Heavy Load & Key Trade-offs
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Heavy Contention & Comparative Trade-offs")
    add_card(s12, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "HEAVY Contention Summary (25 Jobs)", [
        "PSO ranked #1: Score 44.56 +/- 1.13, Makespan 374.51s, SLA violations 56.8%.",
        "GA ranked #2: Score 44.14 +/- 1.12, Makespan 377.56s, SLA violations 56.8%.",
        "SA ranked #3: Score 39.80 +/- 1.08, Makespan 391.92s.",
        "FCFS ranked #4: Score 37.92 (outperformed SJF under heavy load).",
        "SJF dropped to #6: Score 36.49; prioritizing short jobs repeatedly delayed longer batch tasks past deadlines (72.0% SLA breaches).",
        "RL Agents: Q-Learning (23.67) and DQN (20.93) under limited 25–30 episode training budgets."
    ], "Heavy Scenario")
    add_card(s12, Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0), "Core Empirical Takeaways", [
        "1. Low-Load Invariance: In the LIGHT scenario, all heuristics and metaheuristics achieve identical scores (62.49) because resources are never bottlenecked.",
        "2. Metaheuristic Performance: Measured results show PSO and GA achieved the highest efficiency scores; we hypothesize global permutation search mitigates multi-resource fragmentation.",
        "3. SJF Heavy Degradation: SJF achieves strong wait times in moderate load, but suffered 72.0% SLA breaches in heavy load due to long-task postponement.",
        "4. RL Observations: Under a 25–30 episode training budget, RL policies remain suboptimal and do not establish convergence."
    ], "Analysis")
    set_notes(s12, "Under HEAVY contention, measured results show PSO and GA achieved the highest efficiency scores around 44. Crucially, notice SJF: while it ranked 3rd in Medium contention, it dropped to 6th in Heavy contention because repeatedly prioritizing short tasks delayed longer tasks past their deadlines, causing a measured 72.0% SLA breach rate. Under limited 25-30 episode training budgets, RL agents scored 20 to 23.")

    # ==========================================
    # SLIDE 13: Technical Validation & Verification
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "System Verification & Production Readiness")
    add_card(s13, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3), "Backend Test Suite (40/40 Passing)", [
        "Automated regression suite: 40 tests across 9 test files.",
        "Test runtime: 24.04 seconds with zero failures.",
        "Edge-case coverage: empty workloads, solo jobs, simultaneous arrivals (t=0.0), and capacity over-allocation.",
        "Fairness verification: identical deep copies passed to all 10 algorithms."
    ], "Backend QA")
    add_card(s13, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.3), "Frontend Production Build", [
        "Vite production build: Clean compile in 30.32 seconds.",
        "Bundle artifacts: Optimized CSS (22.48 kB) and JS bundle (830.72 kB).",
        "Zero syntax or packaging errors; fully production ready.",
        "Interactive Gantt chart & execution flow sequence verified."
    ], "Frontend QA")
    add_card(s13, Inches(0.8), Inches(4.4), Inches(11.6), Inches(2.4), "Infeasible Job Handling & Zero-Division Safety", [
        "Admissibility Check: Jobs requiring >40 CPU, >256 GB RAM, or >8 GPUs are rejected upfront.",
        "Terminal Status: Explicitly marked status='REJECTED' with detailed diagnostic reason strings.",
        "Zero-Division Protection: Purely infeasible workloads (100% rejected) return 0.0 metrics safely without crashes.",
        "Non-corrupting Semantics: SLA violation rates and compute costs evaluate executed schedules; rejected jobs are reported in a separate audit banner."
    ], "Reliability")
    set_notes(s13, "Our implementation is rigorously verified. We run 40 automated backend tests covering edge cases, algorithmic correctness, and discrepancy detection with 100% pass rates. The frontend builds cleanly via Vite. Furthermore, our engine includes upfront admissibility checks: permanently infeasible jobs are caught immediately, assigned terminal REJECTED status, and prevented from deadlocking simulation queues.")

    # ==========================================
    # SLIDE 14: Limitations, Conclusion & Future Work
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Limitations, Future Enhancements, and Conclusion")
    add_card(s14, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), "Current Limitations", [
        "Non-Preemptive Scheduling: Tasks run uninterrupted; no live checkpointing or task migration.",
        "Aggregated Capacity Pool: Cluster is modeled as a unified capacity pool without NUMA or network topology latency.",
        "RL Training Budgets: Constrained to 25–60 episodes to preserve responsive browser REST interactions."
    ], "Limitations")
    add_card(s14, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), "Future Roadmap", [
        "Preemptive Simulation: Add suspend/resume events to model dynamic priority preemption.",
        "GNN State Representation: Replace 6D summary state vector with bipartite Graph Neural Networks modeling jobs and nodes.",
        "Advanced RL: Implement Proximal Policy Optimization (PPO) with action masking.",
        "Distributed Execution: Celery/Redis background worker queue for long-running benchmarks."
    ], "Future Work")
    add_card(s14, Inches(8.8), Inches(1.8), Inches(3.6), Inches(5.0), "Conclusion", [
        "Successfully engineered a multi-resource scheduling optimization & simulation platform.",
        "Empirically validated 10 algorithms under identical, reproducible conditions.",
        "Demonstrated that PSO and GA achieve superior efficiency under heavy multi-resource contention.",
        "Provided transparent, realistic findings on RL convergence in combinatorial scheduling."
    ], "Summary")
    set_notes(s14, "In conclusion, this project provides a comprehensive, verified platform for multi-resource scheduling optimization. We demonstrated that global metaheuristics like PSO and GA consistently deliver superior efficiency under heavy contention, while classical heuristics like SJF excel only under moderate loads. We have also documented the real-world sample efficiency challenges of reinforcement learning. Thank you, and I look forward to your questions.")

    output_path = os.path.join("docs", "FINAL_PROJECT_PRESENTATION.pptx")
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_deck()

